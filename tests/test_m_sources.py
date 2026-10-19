"""Tier 2 and 3 of pbix_mcp.m: the table functions a query's steps use, and local files (File.Contents,
Csv.Document, Excel.Workbook, Parquet.Document). Rules from Microsoft's function reference; what it leaves open
is refused (each such case is pinned here as an error until Desktop settles it)."""
import datetime as dt
import io
import json
import os
import tempfile
import uuid

import pytest

from pbix_mcp.m import CellError, MError, UnmappedSource, evaluate, force_table


def ev(m, files=None, **names):
    return evaluate(m, files=files, resolve=lambda n: names[n])


def files_of(**contents):
    def read(path):
        if path not in contents:
            raise UnmappedSource(path)
        return contents[path]
    return read


class TestTableFunctions:
    T = '#table({"a", "b"}, {{1, "x-y-z"}, {2, "p-q"}})'

    def test_rename_remove_select(self):
        assert ev(f'Table.RenameColumns({self.T}, {{"a", "A"}})').columns == ["A", "b"]
        assert ev(f'Table.RenameColumns({self.T}, {{{{"a", "A"}}, {{"b", "B"}}}})').columns == ["A", "B"]
        assert ev(f'Table.RemoveColumns({self.T}, "a")').columns == ["b"]
        t = ev(f'Table.SelectColumns({self.T}, {{"b", "a"}})')
        assert t.columns == ["b", "a"] and t.rows[0] == ["x-y-z", 1]

    def test_missing_fields(self):
        with pytest.raises(MError, match="column 'z'"):
            ev(f'Table.RemoveColumns({self.T}, "z")')
        assert ev(f'Table.RemoveColumns({self.T}, "z", MissingField.Ignore)').columns == ["a", "b"]
        with pytest.raises(MError, match="field 'z'"):     # the reference's own example error
            ev(f'Table.SelectColumns({self.T}, {{"a", "z"}})')
        t = ev(f'Table.SelectColumns({self.T}, {{"a", "z"}}, MissingField.UseNull)')
        assert t.columns == ["a", "z"] and t.rows[0] == [1, None]
        assert ev(f'Table.RenameColumns({self.T}, {{"z", "Z"}}, MissingField.Ignore)').columns == ["a", "b"]
        with pytest.raises(MError, match="already exists"):
            ev(f'Table.RenameColumns({self.T}, {{"a", "b"}})')

    def test_add_column(self):
        t = ev(f'Table.AddColumn({self.T}, "c", each [a] * 10, type number)')
        assert t.columns == ["a", "b", "c"] and [r[2] for r in t.rows] == [10, 20] and t.types[2].kind == "number"
        t = ev(f'Table.AddColumn({self.T}, "d", each if [a] = 1 then error "bad" else [a])')
        assert isinstance(t.rows[0][2], CellError) and t.rows[1][2] == 2       # a failed cell keeps its row

    def test_before_delimiter(self):
        # The reference's examples.
        assert ev('Text.BeforeDelimiter("111-222-333", "-")') == "111"
        assert ev('Text.BeforeDelimiter("111-222-333", "-", 1)') == "111-222"
        assert ev('Text.BeforeDelimiter("111-222-333", "-", {1, RelativePosition.FromEnd})') == "111"
        with pytest.raises(MError, match="wasn't found"):
            ev('Text.BeforeDelimiter("111", "-")')

    def test_promote_headers(self):
        t = ev('Table.PromoteHeaders(#table({"Column1", "Column2", "Column3"}, {{"ID", 2024, null}, {1, 2, 3}}))')
        assert t.columns == ["ID", "2024", "Column3"] and t.rows == [[1, 2, 3]]
        with pytest.raises(MError, match="repeats"):
            ev('Table.PromoteHeaders(#table({"Column1", "Column2"}, {{"x", "x"}}))')
        with pytest.raises(MError, match="empty header"):
            ev('Table.PromoteHeaders(#table({"Column1"}, {{""}}))')


class TestFiles:
    def test_without_a_mapping_a_file_is_unmapped(self):
        with pytest.raises(UnmappedSource) as e:
            ev('File.Contents("C:\\data\\x.csv")')
        assert e.value.path == "C:\\data\\x.csv" and e.value.reason == "DataSource.NotFound"

    def test_the_mapping_reads_it(self):
        assert ev('File.Contents("C:\\x")', files_of(**{"C:\\x": b"abc"})) == b"abc"


class TestCsv:
    def csv(self, data, options="", **kw):
        return ev(f'Csv.Document(File.Contents("f"){options})', files_of(f=data))

    def test_quoted_values(self):
        t = self.csv(b'a,"b, c","d ""e"""\r\n1,"two\nlines",3\n')
        assert t.columns == ["Column1", "Column2", "Column3"]
        assert t.rows == [["a", "b, c", 'd "e"'], ["1", "two\nlines", "3"]]

    def test_options_and_columns(self):
        t = self.csv(b"x;y\n1;2\n", ', [Delimiter=";", Columns=2, Encoding=65001, QuoteStyle=QuoteStyle.Csv]')
        assert t.rows == [["x", "y"], ["1", "2"]]
        t = self.csv(b"x,y,z\n1,2,3\n", ', [Columns=2]')            # extra values are ignored by default
        assert t.rows == [["x", "y"], ["1", "2"]]
        with pytest.raises(MError, match="more values"):
            self.csv(b"x,y,z\n", ', [Columns=2, ExtraValues=ExtraValues.Error]')

    def test_encodings_and_byte_order_marks(self):
        assert self.csv("é,ü\n".encode("cp1252"), ', [Encoding=1252]').rows == [["é", "ü"]]
        assert self.csv(b"\xef\xbb\xbfa,b\n").rows == [["a", "b"]]
        assert self.csv(b"\xff\xfe" + "a,b\n".encode("utf-16-le"), ', [Encoding=TextEncoding.Utf16]').rows == [["a", "b"]]

    def test_what_the_reference_leaves_open_is_refused(self):
        with pytest.raises(MError, match="QuoteStyle.None"):
            self.csv(b"a\n", ', [QuoteStyle=QuoteStyle.None]')
        with pytest.raises(MError, match="fewer values"):
            self.csv(b"a,b\n1\n", ', [Columns=2]')
        with pytest.raises(MError, match="different lengths"):
            self.csv(b"a,b\n1\n")


class TestExcel:
    def workbook(self):
        openpyxl = pytest.importorskip("openpyxl")
        from openpyxl.worksheet.table import Table as XTable
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Region", "Amount", "When", "Flag"])
        ws.append(["East", 10.5, dt.datetime(2024, 1, 2), True])
        ws.append(["West", 3, None, False])
        hidden = wb.create_sheet("Hidden")
        hidden["B2"] = "only"
        hidden.sheet_state = "hidden"
        t = wb.create_sheet("T")
        t.append(["Key", "Name"])
        t.append([1, "a"])
        t.append([2, "b"])
        t.add_table(XTable(displayName="Keys", ref="A1:B3"))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_sheets_and_tables(self):
        files = files_of(w=self.workbook())
        nav = ev('Excel.Workbook(File.Contents("w"), null, true)', files)
        assert nav.columns == ["Name", "Data", "Item", "Kind", "Hidden"]
        assert [(r[2], r[3], r[4]) for r in nav.rows] == [("Sales", "Sheet", False), ("Hidden", "Sheet", True),
                                                          ("T", "Sheet", False), ("Keys", "Table", False)]
        sheet = ev('Excel.Workbook(File.Contents("w"), null, true){[Item="Sales",Kind="Sheet"]}[Data]', files)
        assert sheet.columns == ["Column1", "Column2", "Column3", "Column4"]
        assert sheet.rows == [["Region", "Amount", "When", "Flag"], ["East", 10.5, dt.datetime(2024, 1, 2), True],
                              ["West", 3, None, False]]
        table = ev('Excel.Workbook(File.Contents("w"), null, true){[Item="Keys",Kind="Table"]}[Data]', files)
        assert table.columns == ["Key", "Name"] and table.rows == [[1, "a"], [2, "b"]]

    def test_a_sheet_reads_its_stored_dimension(self):
        sheet = ev('Excel.Workbook(File.Contents("w"), null, true){[Item="Hidden",Kind="Sheet"]}[Data]',
                   files_of(w=self.workbook()))
        assert sheet.columns == ["Column1"] and sheet.rows == [["only"]]      # B2:B2

    def test_the_query_desktop_writes(self):
        m = ('let Source = Excel.Workbook(File.Contents("w"), null, true), '
             'Sheet = Source{[Item="Sales",Kind="Sheet"]}[Data], '
             '#"Promoted Headers" = Table.PromoteHeaders(Sheet), '
             '#"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers", '
             '{{"Region", type text}, {"Amount", type number}, {"When", type datetime}, {"Flag", type logical}}) '
             'in #"Changed Type"')
        t = ev(m, files_of(w=self.workbook()))
        assert t.columns == ["Region", "Amount", "When", "Flag"]
        assert force_table(t, errors_as_null=False) == [["East", 10.5, dt.datetime(2024, 1, 2), True],
                                                        ["West", 3, None, False]]

    def test_what_the_reference_leaves_open_is_refused(self):
        files = files_of(w=self.workbook())
        with pytest.raises(MError, match="useHeaders"):
            ev('Excel.Workbook(File.Contents("w"), true, true)', files)
        with pytest.raises(MError, match="delayTypes"):
            ev('Excel.Workbook(File.Contents("w"))', files)


class TestParquet:
    def parquet(self, **columns):
        pa = pytest.importorskip("pyarrow")
        import pyarrow.parquet as pq
        buf = io.BytesIO()
        pq.write_table(pa.table(columns), buf)
        return buf.getvalue()

    def test_plain_types(self):
        pa = pytest.importorskip("pyarrow")
        data = self.parquet(i=pa.array([1, None], pa.int64()), f=pa.array([1.5, 2.0]), s=pa.array(["a", None]),
                            b=pa.array([True, False]), d=pa.array([dt.date(2024, 1, 2), None]),
                            t=pa.array([dt.datetime(2024, 1, 2, 3, 4), None], pa.timestamp("us")))
        t = ev('Parquet.Document(File.Contents("p"))', files_of(p=data))
        assert t.columns == ["i", "f", "s", "b", "d", "t"]
        assert [x.facet or x.kind for x in t.types] == ["Int64.Type", "number", "text", "logical", "date", "datetime"]
        assert t.rows == [[1, 1.5, "a", True, dt.date(2024, 1, 2), dt.datetime(2024, 1, 2, 3, 4)],
                          [None, 2.0, None, False, None, None]]

    def test_other_types_are_refused(self):
        pa = pytest.importorskip("pyarrow")
        with pytest.raises(MError, match="not supported"):
            ev('Parquet.Document(File.Contents("p"))', files_of(p=self.parquet(l=pa.array([[1], [2]]))))


class TestRefreshFromFiles:
    """A table whose M reads a file refreshes through the caller's mapping; without one it keeps its data."""

    def model(self, m):
        from pbix_mcp import server
        from pbix_mcp.builder import PBIXBuilder
        b = PBIXBuilder("Files")
        b.add_table("Sales", [{"name": "Region", "data_type": "String"}, {"name": "Amount", "data_type": "Int64"}],
                    rows=[{"Region": "East", "Amount": 1}])
        b.add_measure("Sales", "Total", "SUM('Sales'[Amount])")
        path = os.path.join(tempfile.mkdtemp(), "files.pbix")
        with open(path, "wb") as f:
            f.write(b.build())
        alias = f"files-{uuid.uuid4().hex[:12]}"
        assert json.loads(server.pbix_open(path, alias))["success"]
        assert json.loads(server.pbix_set_partition_m(alias, "Sales", m))["success"]
        return alias

    def test_mapped_and_unmapped(self):
        from pbix_mcp import server
        from pbix_mcp.refresh import refresh_tables
        m = ('let Source = Csv.Document(File.Contents("C:\\src\\sales.csv"), [Delimiter=",", Columns=2, '
             'Encoding=65001, QuoteStyle=QuoteStyle.Csv]), #"Promoted Headers" = Table.PromoteHeaders(Source), '
             '#"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers", {{"Region", type text}, '
             '{"Amount", Int64.Type}}) in #"Changed Type"')
        alias = self.model(m)
        try:
            result = refresh_tables(alias, ["Sales"])
            assert result.unmapped == {"Sales": "C:\\src\\sales.csv"} and result.tables == {}
            result = refresh_tables(alias, ["Sales"], files=files_of(**{"C:\\src\\sales.csv":
                                                                         b"Region,Amount\nEast,5\nWest,7\n"}))
            assert result.tables == {"Sales": 2} and not result.unmapped
            out = json.loads(server.pbix_evaluate_dax(alias, "Total"))
            assert out["results"][0]["value"] == 12
        finally:
            server.pbix_close(alias, force=True)


class TestReviewFindings:
    """Cases from an independent review of the first version."""

    def test_an_unmapped_file_is_never_absorbed(self):
        with pytest.raises(UnmappedSource):
            ev('try File.Contents("x") otherwise 5')
        with pytest.raises(UnmappedSource):
            ev('Table.AddColumn(#table({"p"}, {{"x"}}), "b", each File.Contents([p]))')

    def test_renames_that_overlap_are_refused(self):
        T = '#table({"a", "b"}, {{1, 2}})'
        for pairs in ('{{"a", "b"}, {"b", "a"}}', '{{"a", "b"}, {"b", "c"}}'):
            with pytest.raises(MError, match="more than one rename"):
                ev(f"Table.RenameColumns({T}, {pairs})")
        assert ev(f'Table.RenameColumns({T}, {{{{"a", "a"}}}})').columns == ["a", "b"]

    def test_csv_malformed_quoting_and_blank_lines_are_refused(self):
        for data in (b'a,"b\n', b'a,"b"c,d\n', b"a\n\nb\n"):
            with pytest.raises(MError):
                ev('Csv.Document(File.Contents("f"))', files_of(f=data))
        assert ev('Csv.Document(File.Contents("f"))', files_of(f=b'a,b"c\r1,2\r')).rows == [['a', 'b"c'], ["1", "2"]]

    def test_csv_options_record_with_positional_arguments_is_refused(self):
        with pytest.raises(MError, match="options record"):
            ev('Csv.Document(File.Contents("f"), [Columns=1], ";")', files_of(f=b"a\n"))

    def xlsx(self, build):
        openpyxl = pytest.importorskip("openpyxl")
        wb = openpyxl.Workbook()
        build(wb)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_excel_honours_the_stored_dimension_for_rows_too(self):
        def build(wb):
            ws = wb.active
            ws.title = "S"
            ws["A1"], ws["C3"] = "x", 3
        data = self.xlsx(build)
        sheet = ev('Excel.Workbook(File.Contents("w"), null, true){[Item="S",Kind="Sheet"]}[Data]', files_of(w=data))
        assert sheet.rows == [["x", None, None], [None, None, None], [None, None, 3]]

    def test_excel_formulas_without_a_cached_value_are_refused(self):
        def build(wb):
            ws = wb.active
            ws.title = "S"
            ws["A1"], ws["A2"] = 1, "=A1+1"      # openpyxl writes no cached value
        with pytest.raises(MError, match="no cached value"):    # Data is read when used
            ev('Excel.Workbook(File.Contents("w"), null, true){[Item="S",Kind="Sheet"]}[Data]',
               files_of(w=self.xlsx(build))).rows

    def test_excel_a_table_with_a_totals_row_is_refused_and_a_table_named_like_a_sheet_is_kept_apart(self):
        from openpyxl.worksheet.table import Table as XTable

        def build(wb):
            ws = wb.active
            ws.title = "Keys"
            ws.append(["K"])
            ws.append([1])
            ws.add_table(XTable(displayName="Keys", ref="A1:A2"))
            t = wb.create_sheet("Tot")
            t.append(["V"])
            t.append([1])
            t.append([1])
            x = XTable(displayName="Tot1", ref="A1:A3")
            x.totalsRowCount = 1
            t.add_table(x)
        files = files_of(w=self.xlsx(build))
        nav = ev('Excel.Workbook(File.Contents("w"), null, true)', files)
        assert [(r[2], r[3]) for r in nav.rows] == [("Keys", "Sheet"), ("Keys", "Table"), ("Tot", "Sheet"),
                                                    ("Tot1", "Table")]
        table = ev('Excel.Workbook(File.Contents("w"), null, true){[Item="Keys",Kind="Table"]}[Data]', files)
        assert table.columns == ["K"] and table.rows == [[1]]
        with pytest.raises(MError, match="totals row"):
            ev('Excel.Workbook(File.Contents("w"), null, true){[Item="Tot1",Kind="Table"]}[Data]', files).rows

    def test_parquet_out_of_range_values_are_refused(self):
        pa = pytest.importorskip("pyarrow")
        import pyarrow.parquet as pq

        def parquet(**cols):
            buf = io.BytesIO()
            pq.write_table(pa.table(cols), buf)
            return buf.getvalue()
        with pytest.raises(MError, match="not supported"):
            ev('Parquet.Document(File.Contents("p"))', files_of(p=parquet(u=pa.array([2 ** 64 - 1], pa.uint64()))))
        ns = pa.array([1_000_000_001], pa.timestamp("ns"))
        with pytest.raises(MError, match="can't be read"):
            ev('Parquet.Document(File.Contents("p"))', files_of(p=parquet(t=ns)))
